"""Export generators for analytics reports (Excel and PDF)."""

import io
from datetime import datetime

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from apps.tokens.analytics import (
    TokenAnalyticsService,
    PhotoValidationAnalyticsService,
    BiometricSubsystemAnalyticsService,
)


class ExcelReportGenerator:
    """Generate Excel reports with multiple sheets and styling."""

    def __init__(self):
        if not OPENPYXL_AVAILABLE:
            raise ImportError('openpyxl is required for Excel report generation')
        self.workbook = None
        self.styles = self._init_styles()

    def _init_styles(self):
        """Initialize reusable cell styles."""
        header_fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )
        
        return {
            'header': {'fill': header_fill, 'font': header_font, 'border': border, 'alignment': Alignment(horizontal='center', vertical='center')},
            'data': {'border': border, 'alignment': Alignment(horizontal='left', vertical='center')},
            'number': {'border': border, 'alignment': Alignment(horizontal='right', vertical='center')},
            'percent': {'border': border, 'alignment': Alignment(horizontal='right', vertical='center'), 'number_format': '0.00"%"'},
        }

    def generate_full_report(self, start_date=None, end_date=None):
        """
        Generate a complete Excel report with multiple sheets.
        
        Returns: BytesIO object with Excel file content
        """
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)  # Remove default sheet

        # Add sheets
        self._add_token_volume_sheet(start_date, end_date)
        self._add_validation_sheet(start_date, end_date)
        self._add_biometric_sheet(start_date, end_date)

        # Write to BytesIO
        output = io.BytesIO()
        self.workbook.save(output)
        output.seek(0)
        return output

    def _add_token_volume_sheet(self, start_date, end_date):
        """Add token volume analytics sheet."""
        ws = self.workbook.create_sheet('Token Volume')
        
        daily_data = TokenAnalyticsService.get_daily_token_volume(start_date, end_date)
        weekly_data = TokenAnalyticsService.get_weekly_token_volume(start_date, end_date)

        # Daily section
        ws.append(['DAILY TOKEN GENERATION VOLUME'])
        ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
        ws['A1'].fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        ws.merge_cells('A1:E1')

        headers = ['Date', 'Total', 'Active', 'Revoked', 'Expired']
        ws.append(headers)
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num)
            cell.value = header
            cell.fill = self.styles['header']['fill']
            cell.font = self.styles['header']['font']
            cell.border = self.styles['header']['border']

        for row_num, data in enumerate(daily_data, 4):
            ws.cell(row=row_num, column=1).value = data['date']
            ws.cell(row=row_num, column=2).value = data['total']
            ws.cell(row=row_num, column=3).value = data.get('ACTIVE', 0)
            ws.cell(row=row_num, column=4).value = data.get('REVOKED', 0)
            ws.cell(row=row_num, column=5).value = data.get('EXPIRED', 0)
            for col in range(1, 6):
                ws.cell(row=row_num, column=col).border = self.styles['data']['border']

        # Weekly section
        weekly_start_row = len(daily_data) + 6
        ws.cell(row=weekly_start_row, column=1).value = 'WEEKLY TOKEN GENERATION VOLUME'
        ws.cell(row=weekly_start_row, column=1).font = Font(bold=True, size=14, color='FFFFFF')
        ws.cell(row=weekly_start_row, column=1).fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        ws.merge_cells(f'A{weekly_start_row}:E{weekly_start_row}')

        headers = ['Week', 'Total', 'Active', 'Revoked', 'Expired']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=weekly_start_row + 1, column=col_num)
            cell.value = header
            cell.fill = self.styles['header']['fill']
            cell.font = self.styles['header']['font']
            cell.border = self.styles['header']['border']

        for row_num, data in enumerate(weekly_data, weekly_start_row + 2):
            ws.cell(row=row_num, column=1).value = data['week']
            ws.cell(row=row_num, column=2).value = data['total']
            ws.cell(row=row_num, column=3).value = data.get('ACTIVE', 0)
            ws.cell(row=row_num, column=4).value = data.get('REVOKED', 0)
            ws.cell(row=row_num, column=5).value = data.get('EXPIRED', 0)
            for col in range(1, 6):
                ws.cell(row=row_num, column=col).border = self.styles['data']['border']

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 12

    def _add_validation_sheet(self, start_date, end_date):
        """Add photo validation analytics sheet."""
        ws = self.workbook.create_sheet('Photo Validation')

        ws.append(['SECURITY PHOTO VALIDATION METRICS'])
        ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
        ws['A1'].fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        ws.merge_cells('A1:C1')

        # Overall metrics
        overall = PhotoValidationAnalyticsService.get_validation_success_rate(start_date=start_date, end_date=end_date)
        metrics_labels = ['Validated', 'Rejected', 'Pending', 'Processing', 'Total', 'Success Rate (%)']
        metrics_values = [
            overall['validated'],
            overall['rejected'],
            overall['pending'],
            overall['processing'],
            overall['total'],
            overall['success_rate'],
        ]

        for row_num, (label, value) in enumerate(zip(metrics_labels, metrics_values), 3):
            ws.cell(row=row_num, column=1).value = label
            ws.cell(row=row_num, column=2).value = value
            ws.cell(row=row_num, column=1).font = Font(bold=True)
            for col in range(1, 3):
                ws.cell(row=row_num, column=col).border = self.styles['data']['border']

        # By event type
        by_event_start_row = len(metrics_labels) + 5
        ws.cell(row=by_event_start_row, column=1).value = 'VALIDATION BY EVENT TYPE'
        ws.cell(row=by_event_start_row, column=1).font = Font(bold=True, size=12, color='FFFFFF')
        ws.cell(row=by_event_start_row, column=1).fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        ws.merge_cells(f'A{by_event_start_row}:D{by_event_start_row}')

        headers = ['Event Type', 'Validated', 'Rejected', 'Success Rate (%)']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=by_event_start_row + 1, column=col_num)
            cell.value = header
            cell.fill = self.styles['header']['fill']
            cell.font = self.styles['header']['font']
            cell.border = self.styles['header']['border']

        by_event = PhotoValidationAnalyticsService.get_validation_by_event_type(start_date=start_date, end_date=end_date)
        for row_num, data in enumerate(by_event, by_event_start_row + 2):
            ws.cell(row=row_num, column=1).value = data['event_type']
            ws.cell(row=row_num, column=2).value = data['validated']
            ws.cell(row=row_num, column=3).value = data['rejected']
            ws.cell(row=row_num, column=4).value = data['success_rate']
            for col in range(1, 5):
                ws.cell(row=row_num, column=col).border = self.styles['data']['border']

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 18

    def _add_biometric_sheet(self, start_date, end_date):
        """Add biometric subsystem analytics sheet."""
        ws = self.workbook.create_sheet('Biometric Status')

        ws.append(['BIOMETRIC SUBSYSTEM PERFORMANCE'])
        ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
        ws['A1'].fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        ws.merge_cells('A1:B1')

        # Status
        status = BiometricSubsystemAnalyticsService.get_face_recognition_status()
        ws.append(['Face Recognition', 'ENABLED' if status['face_recognition_enabled'] else 'DISABLED'])
        ws.append(['Liveness Detection', 'ENABLED' if status['liveness_enabled'] else 'DISABLED'])

        for row in range(3, 5):
            ws.cell(row=row, column=1).font = Font(bold=True)
            for col in range(1, 3):
                ws.cell(row=row, column=col).border = self.styles['data']['border']

        # Performance metrics
        perf_start_row = 6
        ws.cell(row=perf_start_row, column=1).value = 'FACE RECOGNITION PERFORMANCE'
        ws.cell(row=perf_start_row, column=1).font = Font(bold=True, size=12, color='FFFFFF')
        ws.cell(row=perf_start_row, column=1).fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        ws.merge_cells(f'A{perf_start_row}:B{perf_start_row}')

        perf = BiometricSubsystemAnalyticsService.get_face_recognition_performance(start_date=start_date, end_date=end_date)
        perf_labels = ['Success', 'Failed', 'Unavailable', 'Not Run', 'Total', 'Success Rate (%)']
        perf_values = [perf['success'], perf['failed'], perf['unavailable'], perf['not_run'], perf['total'], perf['success_rate']]

        for row_num, (label, value) in enumerate(zip(perf_labels, perf_values), perf_start_row + 1):
            ws.cell(row=row_num, column=1).value = label
            ws.cell(row=row_num, column=2).value = value
            ws.cell(row=row_num, column=1).font = Font(bold=True)
            for col in range(1, 3):
                ws.cell(row=row_num, column=col).border = self.styles['data']['border']

        # Confidence scores
        conf_start_row = perf_start_row + len(perf_labels) + 2
        ws.cell(row=conf_start_row, column=1).value = 'FACE CONFIDENCE STATISTICS'
        ws.cell(row=conf_start_row, column=1).font = Font(bold=True, size=12, color='FFFFFF')
        ws.cell(row=conf_start_row, column=1).fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
        ws.merge_cells(f'A{conf_start_row}:B{conf_start_row}')

        conf = BiometricSubsystemAnalyticsService.get_average_confidence_score(start_date=start_date, end_date=end_date)
        conf_labels = ['Average Confidence', 'Min Confidence', 'Max Confidence', 'Sample Size']
        conf_values = [conf['average_confidence'], conf['min_confidence'], conf['max_confidence'], conf['sample_size']]

        for row_num, (label, value) in enumerate(zip(conf_labels, conf_values), conf_start_row + 1):
            ws.cell(row=row_num, column=1).value = label
            ws.cell(row=row_num, column=2).value = value
            ws.cell(row=row_num, column=1).font = Font(bold=True)
            for col in range(1, 3):
                ws.cell(row=row_num, column=col).border = self.styles['data']['border']

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 15


class PDFReportGenerator:
    """Generate PDF reports using ReportLab."""

    def generate_full_report(self, start_date=None, end_date=None):
        """
        Generate a complete PDF report.
        
        Returns: BytesIO object with PDF file content
        """
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
        story = []
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1F4788'),
            spaceAfter=30,
            alignment=1,  # Center
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#1F4788'),
            spaceAfter=12,
            spaceBefore=12,
        )

        # Title
        story.append(Paragraph('Campus Token Analytics Report', title_style))
        story.append(Spacer(1, 0.2*inch))

        # Token Volume
        story.append(Paragraph('Daily Token Generation Volume', heading_style))
        daily_data = TokenAnalyticsService.get_daily_token_volume(start_date, end_date)
        
        if daily_data:
            table_data = [['Date', 'Total', 'Active', 'Revoked', 'Expired']]
            for data in daily_data[-7:]:  # Last 7 days
                table_data.append([
                    data['date'],
                    str(data['total']),
                    str(data.get('ACTIVE', 0)),
                    str(data.get('REVOKED', 0)),
                    str(data.get('EXPIRED', 0)),
                ])
            
            table = Table(table_data, colWidths=[1.2*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4788')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 11),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.3*inch))

        # Validation Metrics
        story.append(Paragraph('Security Photo Validation Metrics', heading_style))
        validation = PhotoValidationAnalyticsService.get_validation_success_rate(start_date=start_date, end_date=end_date)
        
        validation_data = [
            ['Metric', 'Count'],
            ['Validated', str(validation['validated'])],
            ['Rejected', str(validation['rejected'])],
            ['Pending', str(validation['pending'])],
            ['Processing', str(validation['processing'])],
            ['Total', str(validation['total'])],
            ['Success Rate', f"{validation['success_rate']:.2f}%"],
        ]
        
        table = Table(validation_data, colWidths=[3*inch, 1.5*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        story.append(table)
        story.append(Spacer(1, 0.3*inch))

        # Biometric Status
        story.append(Paragraph('Biometric Subsystem Status', heading_style))
        status = BiometricSubsystemAnalyticsService.get_face_recognition_status()
        biometric = BiometricSubsystemAnalyticsService.get_face_recognition_performance(start_date=start_date, end_date=end_date)
        
        bio_data = [
            ['Status', 'Value'],
            ['Face Recognition', status['status']],
            ['Liveness Detection', 'ENABLED' if status['liveness_enabled'] else 'DISABLED'],
            ['Success Count', str(biometric['success'])],
            ['Failed Count', str(biometric['failed'])],
            ['Success Rate', f"{biometric['success_rate']:.2f}%"],
        ]
        
        table = Table(bio_data, colWidths=[3*inch, 1.5*inch])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ]))
        story.append(table)

        # Build PDF
        doc.build(story)
        output.seek(0)
        return output
